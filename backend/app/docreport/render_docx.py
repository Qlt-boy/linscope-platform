# -*- coding: utf-8 -*-
"""把文档块渲染成 Word（.docx）——python-docx。

版式对齐现场交付的两份样例（服务器 / 交换机健康性检查报告）：

* A4 纵向，上下 2.4cm、左右 2.54cm，正文五号宋体、标题黑体；
* 封面 + 副封面单独一节（不编页码），正文一节从第 1 页起编页码
  （页脚 "第 X 页 / 共 Y 页"）；
* 表格统一 Table Grid + 蓝底白字表头 + 隔行浅底纹，表头跨页重复、行不跨页断开；
* 图片（命令回显截图 / 图表）等比缩放，最高不超过 19cm，避免撑破版面。
"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

from . import fonts

HEAD_FILL = "1A73E8"
ZEBRA_FILL = "F7F9FC"
SUB_COLOR = RGBColor(0x5F, 0x63, 0x68)
INK = RGBColor(0x20, 0x21, 0x24)
TITLE_COLOR = RGBColor(0x14, 0x2A, 0x50)
MAX_IMAGE_HEIGHT_CM = 19.0

ALIGN = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
         "r": WD_ALIGN_PARAGRAPH.RIGHT}


# ---------------------------------------------------------------- 底层工具
def _set_font(run, latin, east, size=None, bold=None, color=None, mono=False):
    if mono:
        latin, east = fonts.WORD_MONO, fonts.WORD_MONO
    run.font.name = latin
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _no_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def _bottom_rule(paragraph, color="1A73E8", size=12):
    """给空段落加一条下边框（当分隔线用）。"""
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    ppr.append(borders)


def _field(paragraph, code, latin, east, size=9):
    run = paragraph.add_run()
    _set_font(run, latin, east, size=size, color=SUB_COLOR)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = code
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _text_run(paragraph, text, *, size=10.5, bold=False, color=None, mono=False,
              latin=None, east=None, italic=False):
    if mono:
        latin, east = fonts.WORD_MONO, fonts.WORD_MONO
    latin = latin or fonts.WORD_BODY[0]
    east = east or fonts.WORD_BODY[1]
    chunks = str(text).split("\n")
    run = None
    for index, chunk in enumerate(chunks):
        if index:
            (run or paragraph.add_run())._r.append(OxmlElement("w:br"))
        run = paragraph.add_run(chunk)
        _set_font(run, latin, east, size=size, bold=bold, color=color, mono=mono)
        run.italic = italic
    return run


def _paragraph(doc, text, *, size=10.5, bold=False, color=None, align=None,
               space_before=0, space_after=6, line_spacing=1.4, indent=None,
               mono=False, latin=None, east=None):
    para = doc.add_paragraph()
    fmt = para.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing
    if indent:
        fmt.left_indent = Cm(indent)
    if align is not None:
        para.alignment = align
    if text:
        _text_run(para, text, size=size, bold=bold, color=color, mono=mono,
                  latin=latin, east=east)
    return para


def _heading(doc, text, level):
    sizes = {1: 15, 2: 13, 3: 12, 4: 11}
    before = {1: 18, 2: 14, 3: 10, 4: 8}
    para = doc.add_paragraph()
    fmt = para.paragraph_format
    fmt.space_before = Pt(before.get(level, 8))
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.3
    _text_run(para, text, size=sizes.get(level, 11), bold=True, color=TITLE_COLOR,
              latin=fonts.WORD_HEAD[0], east=fonts.WORD_HEAD[1])
    return para


# ---------------------------------------------------------------- 图片
def _fit(data: bytes, want_cm: float):
    """按等比缩放给出宽度，同时把高度限制在单页可用范围内。"""
    width_cm = want_cm
    try:
        from PIL import Image
        with Image.open(io.BytesIO(data)) as image:
            ratio = image.height / float(image.width or 1)
        if width_cm * ratio > MAX_IMAGE_HEIGHT_CM:
            width_cm = MAX_IMAGE_HEIGHT_CM / ratio
    except Exception:                                       # noqa: BLE001
        pass
    return width_cm


# ---------------------------------------------------------------- 表格
def _content_width_cm(section) -> float:
    return Emu(section.page_width - section.left_margin - section.right_margin).cm


def _table(doc, block, width_cm):
    headers = block["headers"]
    rows = block["rows"]
    widths = block["widths"]
    aligns = block["align"]
    columns = max([len(headers)] + [len(r) for r in rows] or [1])
    table = doc.add_table(rows=0, cols=columns)
    table.style = doc.styles["Table Grid"]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if not widths:
        widths = [1] * columns
    if len(widths) < columns:
        widths = list(widths) + [1] * (columns - len(widths))
    total = float(sum(widths)) or 1.0
    col_cm = [width_cm * w / total for w in widths[:columns]]

    def fill_row(values, *, head=False, data_index=None):
        row = table.add_row()
        _no_split(row)
        if head:
            _repeat_header(row)
        for index in range(columns):
            cell = row.cells[index]
            cell.width = Cm(col_cm[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(1)
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.line_spacing = 1.15
            align = aligns[index] if index < len(aligns) else "l"
            para.alignment = ALIGN.get("c" if head else align, WD_ALIGN_PARAGRAPH.LEFT)
            value = values[index] if index < len(values) else ""
            _text_run(para, "" if value is None else value,
                      size=9.5 if head else 9, bold=head,
                      color=RGBColor(0xFF, 0xFF, 0xFF) if head else INK)
            if head:
                _shade(cell, HEAD_FILL)
            elif block["zebra"] and data_index is not None and data_index % 2:
                _shade(cell, ZEBRA_FILL)
        return row

    if headers:
        fill_row(headers, head=True)
    for index, values in enumerate(rows):
        fill_row(values, data_index=index)

    for span in block["spans"]:
        r1, c1, r2, c2 = span
        offset = 1 if headers else 0
        try:
            table.cell(r1 + offset, c1).merge(table.cell(r2 + offset, c2))
        except (IndexError, ValueError):
            continue
    for column, width in zip(table.columns, col_cm):
        for cell in column.cells:
            cell.width = Cm(width)
    return table


# ---------------------------------------------------------------- 封面 / 副封面
def _section_geometry(section):
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.4)
    section.left_margin = section.right_margin = Cm(2.54)


def _draw_cover(doc, meta, width_cm):
    cover = meta.get("cover") or {}
    _paragraph(doc, "", space_after=0, line_spacing=1.0)
    _paragraph(doc, "", space_after=0, line_spacing=1.0)
    _paragraph(doc, "", space_after=0, line_spacing=1.0)
    _paragraph(doc, "", space_after=0, line_spacing=1.0)
    _paragraph(doc, cover.get("title") or "", size=24, bold=True, color=TITLE_COLOR,
               align=WD_ALIGN_PARAGRAPH.CENTER, space_before=30, space_after=14,
               line_spacing=1.35, latin=fonts.WORD_TITLE[0], east=fonts.WORD_TITLE[1])
    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rule.paragraph_format.space_after = Pt(30)
    _bottom_rule(rule)
    for _ in range(8):
        _paragraph(doc, "", space_after=0, line_spacing=1.0)
    for line in cover.get("signatures") or []:
        _paragraph(doc, line, size=12, align=WD_ALIGN_PARAGRAPH.RIGHT,
                   space_after=14, line_spacing=1.5)
    for _ in range(7):
        _paragraph(doc, "", space_after=0, line_spacing=1.0)
    _paragraph(doc, cover.get("org") or "", size=13, align=WD_ALIGN_PARAGRAPH.CENTER,
               space_after=2, line_spacing=1.5)
    _paragraph(doc, cover.get("date") or "", size=13, align=WD_ALIGN_PARAGRAPH.CENTER,
               space_after=2, line_spacing=1.5)


def _draw_subtitle_page(doc, meta):
    cover = meta.get("cover") or {}
    lines = cover.get("subtitle_page") or []
    for _ in range(7):
        _paragraph(doc, "", space_after=0, line_spacing=1.0)
    if lines:
        _paragraph(doc, lines[0], size=16, bold=True, color=TITLE_COLOR,
                   space_after=8, line_spacing=1.5, latin=fonts.WORD_HEAD[0],
                   east=fonts.WORD_HEAD[1])
    if len(lines) > 1:
        _paragraph(doc, lines[1], size=14, bold=True, color=TITLE_COLOR,
                   space_after=16, line_spacing=1.5, latin=fonts.WORD_HEAD[0],
                   east=fonts.WORD_HEAD[1])
    for line in lines[2:]:
        _paragraph(doc, line, size=12, space_after=4, line_spacing=1.5)


# ---------------------------------------------------------------- 主入口
def render(report) -> bytes:
    """把 Report 渲染成 .docx 字节。"""
    doc = Document()
    base = doc.styles["Normal"]
    base.font.name = fonts.WORD_BODY[0]
    base.font.size = Pt(10.5)
    try:
        base.element.rPr.rFonts.set(qn("w:eastAsia"), fonts.WORD_BODY[1])
    except AttributeError:                                  # noqa: PERF203
        pass

    section = doc.sections[0]
    _section_geometry(section)
    width_cm = _content_width_cm(section)
    meta = report.meta

    body_section = None
    if meta.get("cover"):
        _draw_cover(doc, meta, width_cm)
        doc.add_page_break()
        _draw_subtitle_page(doc, meta)
        body_section = doc.add_section(WD_SECTION.NEW_PAGE)
        _section_geometry(body_section)

    for kind, payload in report.blocks:
        if kind == "pagebreak":
            doc.add_page_break()
        elif kind == "spacer":
            _paragraph(doc, "", space_after=float(payload or 0), line_spacing=1.0)
        elif kind in ("h1", "h2", "h3", "h4"):
            _heading(doc, payload, int(kind[1]))
        elif kind == "p":
            _paragraph(doc, payload)
        elif kind == "note":
            _paragraph(doc, payload, size=9, color=SUB_COLOR, space_after=4)
        elif kind == "bullets":
            for item in payload:
                _paragraph(doc, item, indent=0.3, space_after=3)
        elif kind == "table":
            _table(doc, payload, width_cm)
            _paragraph(doc, "", space_after=2, line_spacing=1.0)
        elif kind == "img":
            want = min(float(payload.get("width_cm") or width_cm), width_cm)
            width = _fit(payload["data"], want)
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after = Pt(2)
            para.add_run().add_picture(io.BytesIO(payload["data"]), width=Cm(width))
            if payload.get("caption"):
                _paragraph(doc, payload["caption"], size=9, color=SUB_COLOR,
                           align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    if body_section is not None:
        body_section.footer.is_linked_to_previous = False
        para = body_section.footer.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _text_run(para, "第 ", size=9, color=SUB_COLOR)
        _field(para, "PAGE", fonts.WORD_BODY[0], fonts.WORD_BODY[1])
        _text_run(para, " 页 / 共 ", size=9, color=SUB_COLOR)
        _field(para, "NUMPAGES", fonts.WORD_BODY[0], fonts.WORD_BODY[1])
        _text_run(para, " 页", size=9, color=SUB_COLOR)
        pgnum = OxmlElement("w:pgNumType")
        pgnum.set(qn("w:start"), "1")
        body_section._sectPr.append(pgnum)

    core = doc.core_properties
    core.title = report.title
    core.author = meta.get("org") or ""
    core.comments = meta.get("subtitle") or ""

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()
